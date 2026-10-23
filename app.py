import hashlib
import io
import json
import math
import os
import re
import time
import unicodedata
import uuid
from datetime import date, datetime, timedelta
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

import requests
from django.contrib.auth import (
    authenticate,
    login as auth_login,
    logout as auth_logout,
)
from django.contrib.auth.decorators import login_required
from django.db import connection, transaction, close_old_connections
from django.http import HttpResponse, HttpResponseBadRequest
from django.shortcuts import redirect, render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

KST = ZoneInfo("Asia/Seoul")

API_BASE = (
    "https://apis.data.go.kr/1230000/ad/BidPublicInfoService"
)

OPERATIONS = {
    "물품": "getBidPblancListInfoThng",
    "용역": "getBidPblancListInfoServc",
    "공사": "getBidPblancListInfoCnstwk",
    "외자": "getBidPblancListInfoFrgcpt",
}

DEFAULT_CONFIG = {
    "keywords": [
        "시뮬레이션", "simulation", "해석", "CAD", "PLM",
        "Data center", "Datacenter", "데이터센터", "데이터 센터",
    ],
    "exclude": [],
    "mode": "recent",
    "days": 7,
    "start": "",
    "end": "",
    "types": list(OPERATIONS),
    "include_closed": False,
    "auto": True,
}

HEADERS = [
    "구분", "공고종류", "공고번호", "공고차수", "공고명",
    "일치키워드", "공고기관", "수요기관", "공고일시",
    "마감일시", "금액", "금액종류", "공고상태", "원문링크",
]


def now():
    return datetime.now(KST)


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def query(sql, params=()):
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        if cursor.description is None:
            return cursor.rowcount
        names = [column[0] for column in cursor.description]
        return [dict(zip(names, row)) for row in cursor.fetchall()]


def init_tables():
    statements = [
        """
        CREATE TABLE IF NOT EXISTS bw_config (
            id INTEGER PRIMARY KEY,
            value TEXT NOT NULL
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS bw_runs (
            id VARCHAR(100) PRIMARY KEY,
            created TEXT NOT NULL,
            finished TEXT NOT NULL,
            status VARCHAR(20) NOT NULL,
            config TEXT NOT NULL,
            result TEXT NOT NULL,
            error TEXT NOT NULL
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS bw_notices (
            id VARCHAR(200) PRIMARY KEY,
            fingerprint VARCHAR(64) NOT NULL
        )
        """,
    ]
    for sql in statements:
        query(sql)

    query(
        """
        INSERT INTO bw_config (id, value) VALUES (1, %s)
        ON CONFLICT (id) DO NOTHING
        """,
        [dumps(DEFAULT_CONFIG)],
    )


def get_config():
    return json.loads(
        query("SELECT value FROM bw_config WHERE id = 1")[0]["value"]
    )


def split_words(value):
    words = re.split(r"[,，\n]+", value)
    return list(dict.fromkeys(
        word.strip() for word in words if word.strip()
    ))


def normalize(value):
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    return re.sub(r"\s+", "", text)


def matched_keywords(title, config):
    text = normalize(title)

    if any(normalize(word) in text for word in config["exclude"]):
        return []

    return [
        word for word in config["keywords"]
        if normalize(word) in text
    ]


def period(config):
    if config["mode"] == "fixed":
        start = date.fromisoformat(config["start"])
        end = date.fromisoformat(config["end"])
    else:
        end = now().date()
        start = end - timedelta(days=int(config["days"]) - 1)

    if end < start:
        raise ValueError("종료일은 시작일보다 빠를 수 없습니다.")

    if (end - start).days >= 366:
        raise ValueError("1회 조회 기간은 최대 366일입니다.")

    if end > now().date():
        raise ValueError("종료일은 오늘 이후로 지정할 수 없습니다.")

    return start, end


def enqueue(config, run_id=None):
    run_id = run_id or uuid.uuid4().hex

    query(
        """
        INSERT INTO bw_runs
        (id, created, finished, status, config, result, error)
        VALUES (%s, %s, '', '대기', %s, '{}', '')
        ON CONFLICT (id) DO NOTHING
        """,
        [run_id, now().isoformat(), dumps(config)],
    )
    return run_id


def api_items(session, operation, start, end, budget):
    page = 1
    collected = 0

    while True:
        params = {
            "serviceKey": os.environ["DATA_GO_KR_SERVICE_KEY"],
            "type": "json",
            "inqryDiv": "1",
            "inqryBgnDt": start.strftime("%Y%m%d") + "0000",
            "inqryEndDt": end.strftime("%Y%m%d") + "2359",
            "pageNo": page,
            "numOfRows": 100,
        }

        data = None

        for attempt in range(3):
            budget[0] += 1
            if budget[0] > int(os.getenv("MAX_API_REQUESTS", "200")):
                raise RuntimeError(
                    "이번 실행의 API 호출 상한을 초과했습니다. "
                    "조회 기간이나 공고 종류를 줄여 주세요."
                )

            try:
                response = session.get(
                    f"{API_BASE}/{operation}",
                    params=params,
                    timeout=(10, 60),
                )
                response.raise_for_status()
                data = response.json()
                break

            except (requests.RequestException, ValueError):
                if attempt == 2:
                    raise RuntimeError(
                        "API 통신 또는 응답 해석에 실패했습니다. "
                        "인증키·활용승인·서비스 상태를 확인해 주세요."
                    ) from None
                time.sleep(2 ** attempt)

        envelope = data.get("response", {})
        header = envelope.get("header", {})
        code = str(header.get("resultCode", ""))

        if code == "03":
            return

        if code not in {"0", "00", "0000"}:
            raise RuntimeError(
                f"API가 오류를 반환했습니다. 결과코드: {code[:20]}"
            )

        body = envelope.get("body") or {}
        items = body.get("items") or []

        if isinstance(items, dict):
            items = items.get("item", [])

        if isinstance(items, dict):
            items = [items]

        if not isinstance(items, list):
            raise RuntimeError("예상과 다른 API 목록 형식입니다.")

        try:
            total = int(body.get("totalCount", 0))
        except (TypeError, ValueError):
            raise RuntimeError("API 전체 건수를 확인할 수 없습니다.")

        if not items:
            if collected < total:
                raise RuntimeError(
                    "일부 페이지가 비어 있어 수집을 완료하지 못했습니다."
                )
            return

        for item in items:
            if not isinstance(item, dict):
                raise RuntimeError("예상과 다른 공고 데이터 형식입니다.")
            yield item

        collected += len(items)

        if collected >= total:
            return

        page += 1
        time.sleep(0.15)


def is_closed(value):
    if not value:
        return False

    text = str(value).strip()
    try:
        parsed = datetime.fromisoformat(text)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=KST)
        return parsed <= now()
    except ValueError:
        # 해석할 수 없는 마감일은 임의로 제외하지 않습니다.
        return False


def normalize_notice(item, kind, hits):
    amount = ""
    amount_kind = ""

    for field, label in [
        ("bssamt", "기초금액"),
        ("asignBdgtAmt", "배정예산액"),
        ("presmptPrce", "추정가격"),
    ]:
        if item.get(field) not in (None, ""):
            amount = str(item[field])
            amount_kind = label
            break

    return {
        "공고종류": kind,
        "공고번호": str(item.get("bidNtceNo", "")),
        "공고차수": str(item.get("bidNtceOrd", "")),
        "공고명": str(item.get("bidNtceNm", "")),
        "일치키워드": ", ".join(hits),
        "공고기관": str(item.get("ntceInsttNm", "")),
        "수요기관": str(item.get("dminsttNm", "")),
        "공고일시": str(item.get("bidNtceDt", "")),
        "마감일시": str(item.get("bidClseDt", "")),
        "금액": amount,
        "금액종류": amount_kind,
        "공고상태": str(item.get("bidNtceKindNm", "")),
        "원문링크": str(item.get("bidNtceDtlUrl", "")),
    }


def collect(config):
    if not os.getenv("DATA_GO_KR_SERVICE_KEY"):
        raise RuntimeError("서버에 공공데이터포털 인증키가 없습니다.")

    start, end = period(config)
    found = {}
    scanned = 0
    budget = [0]

    with requests.Session() as session:
        for kind in config["types"]:
            chunk_start = start

            while chunk_start <= end:
                chunk_end = min(
                    chunk_start + timedelta(days=6), end
                )

                for item in api_items(
                    session,
                    OPERATIONS[kind],
                    chunk_start,
                    chunk_end,
                    budget,
                ):
                    scanned += 1
                    hits = matched_keywords(
                        item.get("bidNtceNm", ""), config
                    )
                    if not hits:
                        continue

                    if (
                        not config["include_closed"]
                        and is_closed(item.get("bidClseDt"))
                    ):
                        continue

                    row = normalize_notice(item, kind, hits)

                    if not row["공고번호"]:
                        raise RuntimeError(
                            "식별번호가 없는 공고가 반환되었습니다."
                        )

                    key = (
                        f'{kind}:{row["공고번호"]}:{row["공고차수"]}'
                    )
                    found[key] = row

                chunk_start = chunk_end + timedelta(days=1)

    # 전체 수집이 성공한 경우에만 이력을 반영합니다.
    with transaction.atomic():
        for key, row in found.items():
            # 검색어 변경만으로 공고 변경으로 판단하지 않습니다.
            content = {
                key: value for key, value in row.items()
                if key != "일치키워드"
            }
            fingerprint = hashlib.sha256(
                dumps(content).encode("utf-8")
            ).hexdigest()

            old = query(
                "SELECT fingerprint FROM bw_notices WHERE id = %s",
                [key],
            )

            if not old:
                row["구분"] = "신규"
            elif old[0]["fingerprint"] != fingerprint:
                row["구분"] = "변경"
            else:
                row["구분"] = "기존"

            query(
                """
                INSERT INTO bw_notices (id, fingerprint)
                VALUES (%s, %s)
                ON CONFLICT (id)
                DO UPDATE SET fingerprint = excluded.fingerprint
                """,
                [key, fingerprint],
            )

    rows = sorted(
        found.values(),
        key=lambda row: row["공고일시"],
        reverse=True,
    )

    return {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "scanned": scanned,
        "requests": budget[0],
        "rows": rows,
    }


def worker():
    """
    웹 프로세스와 별도로 실행합니다.
    운영 시 worker는 반드시 한 개만 실행합니다.
    """
    print("자동 수집 worker 시작: Asia/Seoul 08:00 / 13:00")

    # 단일 worker 재시작 시 중단 작업을 실패로 표시합니다.
    query(
        """
        UPDATE bw_runs
        SET status = '실패', finished = %s,
            error = '서버 재시작으로 중단되었습니다. 다시 검색해 주세요.'
        WHERE status = '실행중'
        """,
        [now().isoformat()],
    )

    while True:
        close_old_connections()

        try:
            config = get_config()
            current = now()

            if config["auto"]:
                for hour in (8, 13):
                    slot = current.replace(
                        hour=hour, minute=0, second=0, microsecond=0
                    )
                    if current >= slot:
                        # 같은 날짜·시간의 예약 작업은 한 번만 등록됩니다.
                        enqueue(
                            config,
                            f"auto-{current:%Y%m%d}-{hour:02d}",
                        )

            jobs = query(
                """
                SELECT * FROM bw_runs
                WHERE status = '대기'
                ORDER BY created LIMIT 1
                """
            )

            if jobs:
                job = jobs[0]

                claimed = query(
                    """
                    UPDATE bw_runs SET status = '실행중'
                    WHERE id = %s AND status = '대기'
                    """,
                    [job["id"]],
                )

                if claimed:
                    try:
                        result = collect(json.loads(job["config"]))

                        query(
                            """
                            UPDATE bw_runs
                            SET status = '완료', finished = %s,
                                result = %s
                            WHERE id = %s
                            """,
                            [
                                now().isoformat(),
                                dumps(result),
                                job["id"],
                            ],
                        )

                    except Exception as exc:
                        # 요청 URL에는 인증키가 있을 수 있으므로
                        # 예외 원문을 그대로 저장하지 않습니다.
                        message = (
                            str(exc)
                            if isinstance(exc, (RuntimeError, ValueError))
                            else "수집 중 내부 오류가 발생했습니다."
                        )
                        query(
                            """
                            UPDATE bw_runs
                            SET status = '실패', finished = %s,
                                error = %s
                            WHERE id = %s
                            """,
                            [now().isoformat(), message, job["id"]],
                        )

        except Exception:
            # DB 장애가 일시적이면 다음 루프에서 다시 연결합니다.
            print("worker 처리 오류: 연결과 설정을 확인해 주세요.")

        time.sleep(5)


def index(request):
    if not request.user.is_authenticated:
        return render(request, "index.html", {
            "login_only": True,
        })

    jobs = query(
        "SELECT * FROM bw_runs ORDER BY created DESC LIMIT 20"
    )

    selected = request.GET.get("run")
    chosen = next(
        (job for job in jobs if job["id"] == selected),
        jobs[0] if jobs else None,
    )

    result = json.loads(chosen["result"]) if chosen else {}
    config = get_config()

    return render(request, "index.html", {
        "config": config,
        "keywords_text": "\n".join(config["keywords"]),
        "exclude_text": "\n".join(config["exclude"]),
        "types": list(OPERATIONS),
        "jobs": jobs,
        "chosen": chosen,
        "result": result,
        "rows": result.get("rows", []),
    })


@require_POST
def login_view(request):
    user = authenticate(
        request,
        username=request.POST.get("username", ""),
        password=request.POST.get("password", ""),
    )

    if user is None:
        return render(request, "index.html", {
            "login_only": True,
            "error": "아이디 또는 비밀번호를 확인해 주세요.",
        }, status=401)

    auth_login(request, user)
    return redirect("/")


@require_POST
@login_required(login_url="/")
def logout_view(request):
    auth_logout(request)
    return redirect("/")


@require_POST
@login_required(login_url="/")
def save_settings(request):
    try:
        config = {
            "keywords": split_words(
                request.POST.get("keywords", "")
            ),
            "exclude": split_words(
                request.POST.get("exclude", "")
            ),
            "mode": request.POST.get("mode", "recent"),
            "days": int(request.POST.get("days", "7")),
            "start": request.POST.get("start", ""),
            "end": request.POST.get("end", ""),
            "types": [
                kind for kind in request.POST.getlist("types")
                if kind in OPERATIONS
            ],
            "include_closed": "include_closed" in request.POST,
            "auto": "auto" in request.POST,
        }

        if config["mode"] not in {"recent", "fixed"}:
            raise ValueError("잘못된 기간 설정입니다.")

        if not config["keywords"] or not config["types"]:
            raise ValueError("키워드와 공고 종류를 선택해 주세요.")

        if not 1 <= config["days"] <= 366:
            raise ValueError("최근 조회 일수는 1~366일입니다.")

        if len(config["keywords"]) > 100:
            raise ValueError("키워드는 최대 100개입니다.")

        period(config)

    except (ValueError, TypeError):
        return HttpResponseBadRequest(
            "설정을 확인해 주세요. 키워드·공고 종류가 필요하며, "
            "조회 기간은 최대 366일이고 종료일은 오늘까지입니다."
        )

    query(
        "UPDATE bw_config SET value = %s WHERE id = 1",
        [dumps(config)],
    )
    return redirect("/")


@require_POST
@login_required(login_url="/")
def search_now(request):
    # 중복 클릭으로 같은 검색이 쌓이는 것을 방지합니다.
    with transaction.atomic():
        # 설정 행을 잠금 대상으로 사용합니다.
        query(
            "UPDATE bw_config SET value = value WHERE id = 1"
        )
        pending = query(
            """
            SELECT id FROM bw_runs
            WHERE status IN ('대기', '실행중') LIMIT 1
            """
        )
        if pending:
            run_id = pending[0]["id"]
        else:
            run_id = enqueue(get_config())

    return redirect(f"/?run={run_id}")


def make_excel(job):
    result = json.loads(job["result"])
    rows = result.get("rows", [])

    book = Workbook()
    book.remove(book.active)

    selections = [
        ("전체 결과", rows),
        ("신규 공고", [r for r in rows if r["구분"] == "신규"]),
        ("변경 공고", [r for r in rows if r["구분"] == "변경"]),
    ]

    for name, records in selections:
        sheet = book.create_sheet(name)
        sheet.append(HEADERS)

        for record in records:
            sheet.append([
                str(record.get(header, "")) for header in HEADERS
            ])

        # 모든 외부 문자열을 Excel 수식이 아닌 텍스트로 저장합니다.
        for row in sheet.iter_rows():
            for cell in row:
                cell.data_type = "s"

        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill(
                "solid", fgColor="244062"
            )

        for row_number, record in enumerate(records, start=2):
            link = record.get("원문링크", "")
            if urlparse(link).scheme in {"https", "http"}:
                cell = sheet.cell(row_number, len(HEADERS))
                cell.hyperlink = link
                cell.style = "Hyperlink"

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = (
                65 if column[0].value in {"공고명", "원문링크"}
                else 22
            )

    summary = book.create_sheet("실행 요약")
    config = json.loads(job["config"])

    for entry in [
        ("실행 ID", job["id"]),
        ("등록시각", job["created"]),
        ("완료시각", job["finished"]),
        ("조회 시작일", result.get("start", "")),
        ("조회 종료일", result.get("end", "")),
        ("검색 범위", "공고 제목"),
        ("포함 키워드", ", ".join(config["keywords"])),
        ("제외 키워드", ", ".join(config["exclude"])),
        ("조회 원본 건수", result.get("scanned", 0)),
        ("선별 공고 건수", len(rows)),
        ("API 요청 횟수", result.get("requests", 0)),
    ]:
        summary.append([str(value) for value in entry])

    for row in summary.iter_rows():
        for cell in row:
            cell.data_type = "s"

    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 90

    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@login_required(login_url="/")
def download_excel(request, run_id):
    jobs = query(
        "SELECT * FROM bw_runs WHERE id = %s AND status = '완료'",
        [run_id],
    )

    if not jobs:
        return HttpResponse("완료된 결과가 없습니다.", status=404)

    response = HttpResponse(
        make_excel(jobs[0]),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename="bidwatch_{run_id}.xlsx"'
    )
    return response
